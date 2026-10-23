
import os
import io
import re
import json
import time
import asyncio
import logging
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path

from zoneinfo import ZoneInfo

from telegram import (
    Update,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
)
from telegram.constants import ChatAction
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ContextTypes,
    filters,
)

from google import genai
from google.genai import types


# ============================================================
# CONFIGURATION
# ============================================================

TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN", "8962758198:AAFUmzevaYQRmofbMTImrTaXSc9-7tqs8QU")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "AQ.Ab8RN6L5vHywySq2vuq1BML07Fw5xbQCDEpOm1e9jvHse0d5AQ")

# Set the model name to a model available for your API key.
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

TIMEZONE = os.getenv("TIMEZONE", "Europe/Warsaw")
MORNING_HOUR = int(os.getenv("MORNING_HOUR", "9"))

DATABASE_PATH = os.getenv("DATABASE_PATH", "bot_data.sqlite3")

BUFFER_SECONDS = 45
MAX_TELEGRAM_MESSAGE_LENGTH = 4000
MAX_HISTORY_MESSAGES = 12

SUPPORTED_LANGUAGES = {
    "ru": "Русский",
    "en": "English",
    "pl": "Polski",
    "uk": "Українська",
}

DEFAULT_LANGUAGE = "ru"

SYSTEM_PROMPT = """
You are a helpful multilingual educational assistant.

Your responsibilities:
- Explain educational topics clearly.
- Help solve homework problems step by step.
- Explain why an answer is correct.
- Read and analyze text, images, and supported documents.
- Answer in the user's language whenever possible.
- If the user sends a correction, revise your previous answer.
- Never claim to have read a file that you could not access.
- If information is missing, ask a concise clarification question.
- Use readable formatting and mathematical notation where appropriate.
"""


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    level=logging.INFO,
)

logger = logging.getLogger("vakula_bot")


# ============================================================
# GEMINI CLIENT
# ============================================================

gemini_client = (
    genai.Client(api_key=GEMINI_API_KEY)
    if GEMINI_API_KEY
    else None
)


# ============================================================
# DATABASE
# ============================================================

def get_connection():
    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row
    return connection


def init_database():
    with get_connection() as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                language TEXT NOT NULL DEFAULT 'ru',
                created_at TEXT NOT NULL
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)

        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_messages_user_id
            ON messages(user_id, id)
        """)

        connection.commit()


def ensure_user(user_id: int):
    with get_connection() as connection:
        connection.execute(
            """
            INSERT OR IGNORE INTO users
            (user_id, language, created_at)
            VALUES (?, ?, ?)
            """,
            (
                user_id,
                DEFAULT_LANGUAGE,
                datetime.now(ZoneInfo(TIMEZONE)).isoformat(),
            ),
        )
        connection.commit()


def get_user_language(user_id: int) -> str:
    ensure_user(user_id)

    with get_connection() as connection:
        row = connection.execute(
            "SELECT language FROM users WHERE user_id = ?",
            (user_id,),
        ).fetchone()

    if row and row["language"] in SUPPORTED_LANGUAGES:
        return row["language"]

    return DEFAULT_LANGUAGE


def set_user_language(user_id: int, language: str):
    if language not in SUPPORTED_LANGUAGES:
        return

    ensure_user(user_id)

    with get_connection() as connection:
        connection.execute(
            "UPDATE users SET language = ? WHERE user_id = ?",
            (language, user_id),
        )
        connection.commit()


def save_message(user_id: int, role: str, content: str):
    with get_connection() as connection:
        connection.execute(
            """
            INSERT INTO messages
            (user_id, role, content, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (
                user_id,
                role,
                content,
                datetime.now(ZoneInfo(TIMEZONE)).isoformat(),
            ),
        )
        connection.commit()


def get_history(user_id: int, limit: int = MAX_HISTORY_MESSAGES):
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT role, content
            FROM messages
            WHERE user_id = ?
            ORDER BY id DESC
            LIMIT ?
            """,
            (user_id, limit),
        ).fetchall()

    return [
        {"role": row["role"], "content": row["content"]}
        for row in reversed(rows)
    ]


def clear_history(user_id: int):
    with get_connection() as connection:
        connection.execute(
            "DELETE FROM messages WHERE user_id = ?",
            (user_id,),
        )
        connection.commit()


# ============================================================
# TRANSLATIONS
# ============================================================

TEXTS = {
    "ru": {
        "welcome": (
            "Привет! Я учебный помощник.\n\n"
            "Можешь присылать вопросы, фотографии заданий, "
            "голосовые сообщения и документы.\n\n"
            "Команды:\n"
            "/start — главное меню\n"
            "/help — помощь\n"
            "/language — сменить язык\n"
            "/clear — очистить историю\n"
            "/testgemini — проверить подключение к Gemini"
        ),
        "help": (
            "Как пользоваться ботом:\n\n"
            "• Напиши вопрос обычным сообщением.\n"
            "• Отправь фотографию задания.\n"
            "• Отправь голосовое сообщение.\n"
            "• Прикрепи PDF, DOCX или TXT.\n"
            "• Если ответ нужно исправить, напиши замечания.\n\n"
            "Для выбора языка используй /language."
        ),
        "choose_language": "Выбери язык:",
        "language_changed": "Язык изменён на русский.",
        "cleared": "История диалога очищена.",
        "thinking": "Обрабатываю сообщение…",
        "gemini_error": "Не удалось получить ответ от Gemini. Попробуй ещё раз.",
        "missing_config": (
            "Не настроен GEMINI_API_KEY. Добавь его в переменные окружения."
        ),
        "test_ok": "Подключение к Gemini работает.",
        "test_failed": "Проверка Gemini не пройдена. Подробности записаны в журнал.",
        "unsupported_file": "Этот формат файла пока не поддерживается.",
        "empty_message": "Не удалось извлечь текст из сообщения.",
    },
    "en": {
        "welcome": (
            "Hello! I'm your educational assistant.\n\n"
            "Send questions, photos, voice messages, or documents.\n\n"
            "Commands:\n"
            "/start — main menu\n"
            "/help — help\n"
            "/language — change language\n"
            "/clear — clear history\n"
            "/testgemini — test Gemini"
        ),
        "help": (
            "How to use the bot:\n\n"
            "• Send a question.\n"
            "• Send a photo of an exercise.\n"
            "• Send a voice message.\n"
            "• Attach a PDF, DOCX, or TXT file.\n"
            "• Send feedback if an answer needs correction."
        ),
        "choose_language": "Choose a language:",
        "language_changed": "Language changed to English.",
        "cleared": "Conversation history cleared.",
        "thinking": "Processing your message…",
        "gemini_error": "Gemini could not answer. Please try again.",
        "missing_config": "GEMINI_API_KEY is not configured.",
        "test_ok": "Gemini connection is working.",
        "test_failed": "Gemini test failed. See the logs for details.",
        "unsupported_file": "This file format is not supported yet.",
        "empty_message": "Could not extract any text from the message.",
    },
    "pl": {
        "welcome": (
            "Cześć! Jestem asystentem edukacyjnym.\n\n"
            "Wysyłaj pytania, zdjęcia, wiadomości głosowe i dokumenty.\n\n"
            "/start — menu główne\n"
            "/help — pomoc\n"
            "/language — zmiana języka\n"
            "/clear — wyczyść historię\n"
            "/testgemini — sprawdź Gemini"
        ),
        "help": (
            "Jak korzystać z bota:\n\n"
            "• Wyślij pytanie.\n"
            "• Wyślij zdjęcie zadania.\n"
            "• Wyślij wiadomość głosową.\n"
            "• Załącz PDF, DOCX lub TXT.\n"
            "• Prześlij uwagi, jeśli odpowiedź wymaga poprawy."
        ),
        "choose_language": "Wybierz język:",
        "language_changed": "Zmieniono język na polski.",
        "cleared": "Historia rozmowy została wyczyszczona.",
        "thinking": "Przetwarzam wiadomość…",
        "gemini_error": "Nie udało się uzyskać odpowiedzi od Gemini.",
        "missing_config": "Nie skonfigurowano GEMINI_API_KEY.",
        "test_ok": "Połączenie z Gemini działa.",
        "test_failed": "Test Gemini nie powiódł się. Sprawdź logi.",
        "unsupported_file": "Ten format pliku nie jest obsługiwany.",
        "empty_message": "Nie udało się odczytać tekstu z wiadomości.",
    },
    "uk": {
        "welcome": (
            "Привіт! Я навчальний помічник.\n\n"
            "Надсилай запитання, фотографії, голосові повідомлення "
            "та документи.\n\n"
            "/start — головне меню\n"
            "/help — допомога\n"
            "/language — змінити мову\n"
            "/clear — очистити історію\n"
            "/testgemini — перевірити Gemini"
        ),
        "help": (
            "Як користуватися ботом:\n\n"
            "• Надішли запитання.\n"
            "• Надішли фото завдання.\n"
            "• Надішли голосове повідомлення.\n"
            "• Прикріпи PDF, DOCX або TXT.\n"
            "• Надішли зауваження, якщо відповідь потрібно виправити."
        ),
        "choose_language": "Обери мову:",
        "language_changed": "Мову змінено на українську.",
        "cleared": "Історію діалогу очищено.",
        "thinking": "Обробляю повідомлення…",
        "gemini_error": "Не вдалося отримати відповідь від Gemini.",
        "missing_config": "Не налаштовано GEMINI_API_KEY.",
        "test_ok": "З'єднання з Gemini працює.",
        "test_failed": "Перевірка Gemini не пройшла. Переглянь журнал.",
        "unsupported_file": "Цей формат файлу поки не підтримується.",
        "empty_message": "Не вдалося отримати текст із повідомлення.",
    },
}


def tr(user_id: int, key: str) -> str:
    language = get_user_language(user_id)
    return TEXTS.get(language, TEXTS["ru"]).get(
        key,
        TEXTS["ru"].get(key, key),
    )


# ============================================================
# MESSAGE BUFFER
# ============================================================

message_buffers = {}
buffer_tasks = {}
buffer_locks = {}


def get_user_lock(user_id: int) -> asyncio.Lock:
    if user_id not in buffer_locks:
        buffer_locks[user_id] = asyncio.Lock()

    return buffer_locks[user_id]


async def add_to_buffer(user_id: int, content):
    if user_id not in message_buffers:
        message_buffers[user_id] = []

    message_buffers[user_id].append(content)

    old_task = buffer_tasks.get(user_id)

    if old_task and not old_task.done():
        old_task.cancel()

    buffer_tasks[user_id] = asyncio.create_task(
        process_buffer_after_delay(user_id)
    )


async def process_buffer_after_delay(user_id: int):
    try:
        await asyncio.sleep(BUFFER_SECONDS)

        async with get_user_lock(user_id):
            items = message_buffers.pop(user_id, [])

            if not items:
                return

            await process_user_content(user_id, items)

    except asyncio.CancelledError:
        return
    except Exception:
        logger.exception("Buffer processing failed for user %s", user_id)


# ============================================================
# GEMINI REQUEST
# ============================================================

def make_history_contents(history):
    contents = []

    for item in history:
        role = "user" if item["role"] == "user" else "model"

        contents.append(
            types.Content(
                role=role,
                parts=[types.Part.from_text(text=item["content"])],
            )
        )

    return contents


async def ask_gemini(user_id: int, user_parts, prompt_text: str = ""):
    if gemini_client is None:
        raise RuntimeError("GEMINI_API_KEY is not configured")

    history = get_history(user_id)

    contents = make_history_contents(history)

    current_parts = []

    if prompt_text.strip():
        current_parts.append(
            types.Part.from_text(text=prompt_text.strip())
        )

    current_parts.extend(user_parts)

    if not current_parts:
        current_parts.append(
            types.Part.from_text(text="Please help me with this message.")
        )

    contents.append(
        types.Content(
            role="user",
            parts=current_parts,
        )
    )

    response = await asyncio.to_thread(
        gemini_client.models.generate_content,
        model=GEMINI_MODEL,
        contents=contents,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            temperature=0.4,
        ),
    )

    answer = getattr(response, "text", None)

    if not answer:
        raise RuntimeError("Gemini returned an empty response")

    user_text = prompt_text.strip() or "[Media or document message]"

    save_message(user_id, "user", user_text)
    save_message(user_id, "assistant", answer)

    return answer


# ============================================================
# TELEGRAM RESPONSE HELPERS
# ============================================================

async def send_long_message(message, text: str):
    if not text:
        text = "Пустой ответ."

    for index in range(0, len(text), MAX_TELEGRAM_MESSAGE_LENGTH):
        chunk = text[index:index + MAX_TELEGRAM_MESSAGE_LENGTH]

        await message.reply_text(
            chunk,
            disable_web_page_preview=True,
        )


async def send_error(message, user_id: int, exc: Exception):
    logger.exception(
        "Error while processing user %s",
        user_id,
        exc_info=exc,
    )

    await message.reply_text(tr(user_id, "gemini_error"))


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension in {".txt", ".md", ".csv", ".json", ".py"}:
        return file_bytes.decode("utf-8", errors="replace")

    if extension == ".docx":
        try:
            from docx import Document

            document = Document(io.BytesIO(file_bytes))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            return "\n".join(paragraphs)

        except Exception as exc:
            logger.warning("DOCX extraction failed: %s", exc)
            return ""

    if extension == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(file_bytes))

            pages = []

            for page in reader.pages:
                pages.append(page.extract_text() or "")

            return "\n\n".join(pages)

        except Exception as exc:
            logger.warning("PDF extraction failed: %s", exc)
            return ""

    return ""


# ============================================================
# PROCESS USER CONTENT
# ============================================================

async def process_user_content(user_id: int, items):
    if not items:
        return

    # The bot uses the most recent update's chat context, saved
    # in pending_contexts when the message enters the buffer.
    context = pending_contexts.get(user_id)

    if context is None:
        return

    message = context["message"]
    telegram_context = context["context"]

    await telegram_context.bot.send_chat_action(
        chat_id=message.chat_id,
        action=ChatAction.TYPING,
    )

    prompt_parts = []
    media_parts = []

    for item in items:
        item_type = item.get("type")

        if item_type == "text":
            prompt_parts.append(item["text"])

        elif item_type == "image":
            media_parts.append(
                types.Part.from_bytes(
                    data=item["data"],
                    mime_type=item.get("mime_type", "image/jpeg"),
                )
            )

        elif item_type == "document":
            extracted = item.get("text", "")

            if extracted:
                prompt_parts.append(
                    "Document: "
                    + item.get("filename", "document")
                    + "\n\n"
                    + extracted
                )

            elif item.get("mime_type") == "application/pdf":
                media_parts.append(
                    types.Part.from_bytes(
                        data=item["data"],
                        mime_type="application/pdf",
                    )
                )

        elif item_type == "audio":
            media_parts.append(
                types.Part.from_bytes(
                    data=item["data"],
                    mime_type=item.get("mime_type", "audio/ogg"),
                )
            )
            prompt_parts.append(
                "Please transcribe and respond to the audio message."
            )

    prompt = "\n\n".join(
        text for text in prompt_parts if text.strip()
    )

    if not prompt and not media_parts:
        await message.reply_text(tr(user_id, "empty_message"))
        return

    try:
        answer = await ask_gemini(
            user_id=user_id,
            user_parts=media_parts,
            prompt_text=prompt,
        )

        await send_long_message(message, answer)

    except Exception as exc:
        logger.exception("Gemini request failed")

        if "not configured" in str(exc).lower():
            await message.reply_text(tr(user_id, "missing_config"))
        else:
            await message.reply_text(tr(user_id, "gemini_error"))


pending_contexts = {}


# ============================================================
# COMMANDS
# ============================================================

async def start_command(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message:
        return

    ensure_user(user.id)

    await message.reply_text(tr(user.id, "welcome"))


async def help_command(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if user and message:
        await message.reply_text(tr(user.id, "help"))


async def clear_command(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message:
        return

    clear_history(user.id)
    message_buffers.pop(user.id, None)

    task = buffer_tasks.pop(user.id, None)

    if task and not task.done():
        task.cancel()

    await message.reply_text(tr(user.id, "cleared"))


async def language_command(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message:
        return

    keyboard = [
        [
            InlineKeyboardButton(
                "🇷🇺 Русский",
                callback_data="lang:ru",
            ),
            InlineKeyboardButton(
                "🇬🇧 English",
                callback_data="lang:en",
            ),
        ],
        [
            InlineKeyboardButton(
                "🇵🇱 Polski",
                callback_data="lang:pl",
            ),
            InlineKeyboardButton(
                "🇺🇦 Українська",
                callback_data="lang:uk",
            ),
        ],
    ]

    await message.reply_text(
        tr(user.id, "choose_language"),
        reply_markup=InlineKeyboardMarkup(keyboard),
    )


async def language_callback(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    query = update.callback_query

    if not query or not query.from_user:
        return

    await query.answer()

    data = query.data or ""

    if not data.startswith("lang:"):
        return

    language = data.split(":", 1)[1]

    if language not in SUPPORTED_LANGUAGES:
        return

    set_user_language(query.from_user.id, language)

    try:
        await query.edit_message_text(
            TEXTS[language]["language_changed"]
        )
    except Exception:
        logger.exception("Could not edit language message")


async def test_gemini_command(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message:
        return

    if gemini_client is None:
        await message.reply_text(tr(user.id, "missing_config"))
        return

    try:
        response = await asyncio.to_thread(
            gemini_client.models.generate_content,
            model=GEMINI_MODEL,
            contents="Reply with exactly: OK",
        )

        if getattr(response, "text", None):
            await message.reply_text(tr(user.id, "test_ok"))
        else:
            await message.reply_text(tr(user.id, "test_failed"))

    except Exception:
        logger.exception("Gemini connectivity test failed")
        await message.reply_text(tr(user.id, "test_failed"))


# ============================================================
# MESSAGE HANDLERS
# ============================================================

async def handle_text(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message or not message.text:
        return

    ensure_user(user.id)

    pending_contexts[user.id] = {
        "message": message,
        "context": context,
    }

    await add_to_buffer(
        user.id,
        {
            "type": "text",
            "text": message.text,
        },
    )


async def handle_photo(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message or not message.photo:
        return

    ensure_user(user.id)

    photo = message.photo[-1]
    telegram_file = await context.bot.get_file(photo.file_id)

    buffer = io.BytesIO()
    await telegram_file.download_to_memory(out=buffer)

    pending_contexts[user.id] = {
        "message": message,
        "context": context,
    }

    await add_to_buffer(
        user.id,
        {
            "type": "image",
            "data": buffer.getvalue(),
            "mime_type": "image/jpeg",
        },
    )


async def handle_document(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message or not message.document:
        return

    ensure_user(user.id)

    document = message.document
    filename = document.file_name or "document"

    allowed_extensions = {
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".py",
        ".pdf",
        ".docx",
    }

    extension = Path(filename).suffix.lower()

    if extension not in allowed_extensions:
        await message.reply_text(tr(user.id, "unsupported_file"))
        return

    telegram_file = await context.bot.get_file(document.file_id)

    buffer = io.BytesIO()
    await telegram_file.download_to_memory(out=buffer)

    file_bytes = buffer.getvalue()

    extracted_text = extract_text_from_file(
        filename,
        file_bytes,
    )

    mime_type = document.mime_type or ""

    pending_contexts[user.id] = {
        "message": message,
        "context": context,
    }

    await add_to_buffer(
        user.id,
        {
            "type": "document",
            "filename": filename,
            "text": extracted_text,
            "data": file_bytes,
            "mime_type": mime_type,
        },
    )


async def handle_voice(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    user = update.effective_user
    message = update.effective_message

    if not user or not message:
        return

    voice = message.voice

    if voice is None:
        voice = message.audio

    if voice is None:
        return

    ensure_user(user.id)

    telegram_file = await context.bot.get_file(voice.file_id)

    buffer = io.BytesIO()
    await telegram_file.download_to_memory(out=buffer)

    mime_type = getattr(voice, "mime_type", None) or "audio/ogg"

    pending_contexts[user.id] = {
        "message": message,
        "context": context,
    }

    await add_to_buffer(
        user.id,
        {
            "type": "audio",
            "data": buffer.getvalue(),
            "mime_type": mime_type,
        },
    )


# ============================================================
# MORNING GREETING
# ============================================================

async def morning_job(
    context: ContextTypes.DEFAULT_TYPE,
):
    now = datetime.now(ZoneInfo(TIMEZONE))
    logger.info("Morning job executed at %s", now.isoformat())

    # Telegram bots cannot message users unless their chat ID
    # is known. This job sends a greeting to users who have
    # previously interacted with the bot.
    with get_connection() as connection:
        rows = connection.execute(
            "SELECT user_id, language FROM users"
        ).fetchall()

    greetings = {
        "ru": "Доброе утро! ☀️ Хорошего дня и продуктивной учёбы!",
        "en": "Good morning! ☀️ Have a great and productive day!",
        "pl": "Dzień dobry! ☀️ Miłego i produktywnego dnia!",
        "uk": "Доброго ранку! ☀️ Гарного та продуктивного дня!",
    }

    for row in rows:
        user_id = row["user_id"]
        language = row["language"]

        try:
            await context.bot.send_message(
                chat_id=user_id,
                text=greetings.get(language, greetings["ru"]),
            )
        except Exception:
            logger.warning(
                "Could not send morning greeting to user %s",
                user_id,
            )


# ============================================================
# ERROR HANDLER
# ============================================================

async def error_handler(
    update: object,
    context: ContextTypes.DEFAULT_TYPE,
):
    logger.error(
        "Unhandled Telegram error",
        exc_info=context.error,
    )


# ============================================================
# APPLICATION
# ============================================================

def main():
    if not TELEGRAM_TOKEN:
        raise RuntimeError(
            "TELEGRAM_TOKEN is missing. Set it in environment variables."
        )

    init_database()

    application = (
        Application.builder()
        .token(TELEGRAM_TOKEN)
        .build()
    )

    application.add_handler(CommandHandler("start", start_command))
    application.add_handler(CommandHandler("help", help_command))
    application.add_handler(CommandHandler("clear", clear_command))
    application.add_handler(CommandHandler("language", language_command))
    application.add_handler(
        CommandHandler("testgemini", test_gemini_command)
    )

    application.add_handler(
        CallbackQueryHandler(
            language_callback,
            pattern=r"^lang:",
        )
    )

    application.add_handler(
        MessageHandler(filters.PHOTO, handle_photo)
    )

    application.add_handler(
        MessageHandler(
            filters.Document.ALL,
            handle_document,
        )
    )

    application.add_handler(
        MessageHandler(
            filters.VOICE | filters.AUDIO,
            handle_voice,
        )
    )

    application.add_handler(
        MessageHandler(
            filters.TEXT & ~filters.COMMAND,
            handle_text,
        )
    )

    application.add_error_handler(error_handler)

    if application.job_queue is None:
        logger.warning(
            "JobQueue is unavailable. Install "
            "python-telegram-bot[job-queue] to enable morning greetings."
        )
    else:
        application.job_queue.run_daily(
            morning_job,
            time=datetime.strptime(
                f"{MORNING_HOUR}:00",
                "%H:%M",
            ).time().replace(tzinfo=ZoneInfo(TIMEZONE)),
            name="morning_greeting",
        )

    logger.info("Starting Vakula educational bot")

    application.run_polling(
        allowed_updates=Update.ALL_TYPES,
        drop_pending_updates=False,
    )


if __name__ == "__main__":
    main()
